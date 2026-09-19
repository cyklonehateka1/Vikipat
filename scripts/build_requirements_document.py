from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parents[1] / "docs" / "Vikipat_Platform_Discovery_and_Service_Pricing_Workbook.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = "08194A"; BLUE = "1234D8"; CYAN = "0CA8E8"; LIME = "B8E61D"
INK = "14213D"; MUTED = "56627A"; PALE = "EDF3FF"; LIGHT = "F5F7FB"; WHITE = "FFFFFF"; BORDER = "CAD3E1"
ACCENTS = ["F6B21A", "10A96B", "E83E8C", "0CA8E8"]

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5); sec.page_height = Inches(11)
sec.top_margin = Inches(.72); sec.bottom_margin = Inches(.7); sec.left_margin = Inches(.78); sec.right_margin = Inches(.78)
sec.header_distance = Inches(.32); sec.footer_distance = Inches(.35)

def font(run, name="Aptos", size=10.5, bold=False, color=INK, italic=False):
    run.font.name = name; run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name); run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size); run.bold = bold; run.italic = italic; run.font.color.rgb = RGBColor.from_string(color)

styles = doc.styles
normal = styles["Normal"]; normal.font.name = "Aptos"; normal.font.size = Pt(10.5); normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.18
for name, size, color, before, after in [("Title",30,NAVY,0,8),("Subtitle",13,MUTED,0,18),("Heading 1",18,NAVY,18,8),("Heading 2",13.5,BLUE,13,6),("Heading 3",11.5,NAVY,10,4)]:
    st = styles[name]; st.font.name="Aptos Display" if name != "Normal" else "Aptos"; st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color); st.font.bold=name!="Subtitle"
    st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn("w:shd")) or OxmlElement("w:shd"); shd.set(qn("w:fill"),fill)
    if shd.getparent() is None: tcPr.append(shd)

def margins(cell, top=100, start=120, bottom=100, end=120):
    tc=cell._tc.get_or_add_tcPr(); tcMar=tc.first_child_found_in("w:tcMar")
    if tcMar is None: tcMar=OxmlElement("w:tcMar"); tc.append(tcMar)
    for side,val in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node=tcMar.find(qn(f"w:{side}")) or OxmlElement(f"w:{side}"); node.set(qn("w:w"),str(val)); node.set(qn("w:type"),"dxa")
        if node.getparent() is None: tcMar.append(node)

def set_width(cell, dxa):
    tcPr=cell._tc.get_or_add_tcPr(); tcW=tcPr.find(qn("w:tcW")) or OxmlElement("w:tcW"); tcW.set(qn("w:w"),str(dxa)); tcW.set(qn("w:type"),"dxa")
    if tcW.getparent() is None: tcPr.append(tcW)

def table(headers, rows, widths=None, font_size=9):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    tPr=t._tbl.tblPr; tblW=tPr.find(qn("w:tblW")); tblW.set(qn("w:w"),"9940"); tblW.set(qn("w:type"),"dxa")
    ind=OxmlElement("w:tblInd"); ind.set(qn("w:w"),"120"); ind.set(qn("w:type"),"dxa"); tPr.append(ind)
    widths=widths or [9940//len(headers)]*len(headers)
    for i,(c,h) in enumerate(zip(t.rows[0].cells,headers)):
        shade(c,NAVY); margins(c,120,130,120,130); set_width(c,widths[i]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); font(p.add_run(h),size=font_size,bold=True,color=WHITE)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,(c,v) in enumerate(zip(cells,row)):
            set_width(c,widths[i]); margins(c,110,130,110,130); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2: shade(c,LIGHT)
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); font(p.add_run(str(v)),size=font_size,color=INK)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def p(text="", bold=False, color=INK, size=10.5, align=None, italic=False, after=6):
    para=doc.add_paragraph(); para.paragraph_format.space_after=Pt(after)
    if align is not None: para.alignment=align
    font(para.add_run(text),size=size,bold=bold,color=color,italic=italic); return para

def bullet(text, level=0):
    para=doc.add_paragraph(style="List Bullet" if level==0 else "List Bullet 2"); para.paragraph_format.space_after=Pt(3); font(para.add_run(text),size=10.2); return para

def check(text):
    para=doc.add_paragraph(); para.paragraph_format.left_indent=Inches(.18); para.paragraph_format.first_line_indent=Inches(-.18); para.paragraph_format.space_after=Pt(3)
    font(para.add_run("☐ "),size=11,bold=True,color=BLUE); font(para.add_run(text),size=10.2); return para

def field(label, hint="", lines=1):
    para=doc.add_paragraph(); para.paragraph_format.space_after=Pt(3); font(para.add_run(label+": "),bold=True,color=NAVY)
    if hint: font(para.add_run(hint),size=9.5,color=MUTED,italic=True)
    for _ in range(lines): p("________________________________________________________________________________",color=BORDER,size=9,after=3)

def callout(title,text,fill=PALE):
    t=doc.add_table(rows=1,cols=1); t.autofit=False; t.alignment=WD_TABLE_ALIGNMENT.CENTER; c=t.cell(0,0); set_width(c,9940); margins(c,150,180,150,180); shade(c,fill)
    pp=c.paragraphs[0]; pp.paragraph_format.space_after=Pt(3); font(pp.add_run(title),bold=True,color=NAVY,size=10.5)
    pp=c.add_paragraph(); pp.paragraph_format.space_after=Pt(0); font(pp.add_run(text),size=9.8,color=INK)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def section(title, intro=None):
    doc.add_heading(title,level=1)
    if intro: p(intro,color=MUTED,size=10)

def subsection(title,intro=None):
    doc.add_heading(title,level=2)
    if intro: p(intro,color=MUTED,size=9.8)

def page_break(): doc.add_page_break()

# Running header and footer
hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; font(hp.add_run("VIKIPAT  /  PLATFORM DISCOVERY WORKBOOK"),size=8,bold=True,color=MUTED)
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(fp.add_run("Confidential • Prepared for Vikipat Media Solutions • "),size=8,color=MUTED)
fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),"PAGE"); fp._p.append(fld)

# Cover
p("VIKIPAT",bold=True,color=BLUE,size=13,align=WD_ALIGN_PARAGRAPH.CENTER,after=54)
p("Digital Commerce &\nProduction Operations Platform",bold=True,color=NAVY,size=30,align=WD_ALIGN_PARAGRAPH.CENTER,after=12)
p("Discovery Questionnaire, Service Catalogue & Pricing Workbook",color=BLUE,size=15,align=WD_ALIGN_PARAGRAPH.CENTER,after=34)
callout("Purpose", "This workbook captures the information required to design a full customer-ordering, intelligent pricing, Paystack payment, staff production, fulfilment and executive financial-management system for Vikipat.")
p("Prepared for",bold=True,color=MUTED,size=9,align=WD_ALIGN_PARAGRAPH.CENTER,after=3)
p("Vikipat Media Solutions",bold=True,color=NAVY,size=15,align=WD_ALIGN_PARAGRAPH.CENTER,after=2)
p("Mallam–Gbawe Road, opposite Zen Filling Station, Accra",color=MUTED,size=10,align=WD_ALIGN_PARAGRAPH.CENTER,after=30)
p("Completion status",bold=True,color=MUTED,size=9,align=WD_ALIGN_PARAGRAPH.CENTER,after=8)
table(["Business owner review","Operations review","Finance review","Final approval"],[["☐ Pending","☐ Pending","☐ Pending","☐ Pending"]],[2485]*4,9)
p("Version 1.0  •  5 September 2026",color=MUTED,size=9,align=WD_ALIGN_PARAGRAPH.CENTER,after=0)
page_break()

section("How to use this workbook")
p("Complete this workbook with the people who actually quote jobs, operate the machines, buy materials, approve discounts, receive payments and dispatch finished work. Attach current price lists, sample invoices, job cards and artwork specifications wherever requested.")
for x in ["Use one Service Pricing Worksheet for every distinct service or product family.","Provide real examples: at least five recent jobs per major service, showing how the final price was reached.","Mark information that changes frequently and identify who may update it in the system.","Do not combine costs that are calculated differently. Separate design, setup, material, printing, finishing, delivery and taxes.","Agree which prices customers may see instantly and which jobs must be reviewed by a human."] : check(x)
callout("Critical outcome", "The platform can only calculate prices accurately when every price driver, minimum charge, wastage rule, quantity break, machine constraint and exception is documented.", "FFF4D8")
subsection("People who should participate")
table(["Area","Required contributor","Name / role","Sign-off"],[["Leadership","Owner / managing director","","☐"],["Sales","Quoting or customer-service lead","","☐"],["Production","Large-format, digital/offset and embroidery leads","","☐"],["Finance","Accountant / cashier","","☐"],["Procurement","Materials and supplier lead","","☐"],["Delivery","Dispatch / logistics lead","","☐"]],[1800,3000,3500,1640],9)

section("1. Business identity, scope and operating model")
for label,hint in [("Registered business name","Exact legal name"),("Trading name","Name customers see"),("Registration / tax identifiers","TIN, VAT status and registration numbers"),("Primary email","Orders, receipts and system notifications"),("Primary phone numbers","Include country code"),("WhatsApp number","Include country code"),("Workshop address","Landmark and digital address"),("Other branches / pickup points","Address and opening hours"),("Operating hours","Weekdays, weekends and holidays"),("Primary system owner","Name, role, email and phone")]: field(label,hint)
subsection("Business channels and customers")
for x in ["Walk-in retail customers","Online retail customers","SMEs and growing brands","Corporate clients","Schools and institutions","Government / public sector","Event organisers","Resellers / print brokers","International customers"]: check(x)
field("Typical monthly order count","Approximate orders by channel")
field("Average order value","Low, typical and high values")
field("Seasonal peaks","Months, events and affected services")
field("Top ten customers or customer types","Do not include confidential figures unless approved",2)

section("2. Platform vision and success measures")
field("What should customers be able to complete without staff help?","Describe ideal self-service journey",2)
field("What must always be reviewed by staff before payment?","Complex, risky or bespoke jobs",2)
field("Biggest problems with today’s process","Quoting, revisions, payments, job tracking, waste, delivery, reporting",3)
field("Success after 90 days","Measurable outcomes",2)
field("Success after 12 months","Revenue, speed, accuracy, customer satisfaction, capacity",2)
table(["Metric","Current baseline","12-month target","How measured"],[["Quote response time","","",""],["Order-to-production time","","",""],["On-time delivery rate","","",""],["Pricing error rate","","",""],["Gross margin","","",""],["Repeat purchase rate","","",""]],[2700,2200,2200,2840],8.6)

page_break(); section("3. Complete service catalogue")
p("List every service sold today, services planned within 12 months, and services subcontracted to partners. Each selected item must receive its own detailed pricing worksheet later in this document.")
service_groups = {
"Design & prepress":["Logo design","Brand identity systems","Graphic design","Artwork recreation / vectorisation","Layout and typesetting","Photo editing / retouching","Preflight / file correction","Mock-ups and proofs"],
"Small-format printing":["Business cards","Flyers / handbills","Brochures","Posters","Letterheads","Envelopes","Receipt / invoice books","Calendars","Certificates","Books / manuals","Menus","Invitation cards","Funeral programmes","Stickers / labels"],
"Large-format printing":["PVC banners","Pull-up banners","Backdrops","Billboards","Vinyl stickers","One-way vision","Wall graphics","Floor graphics","Canvas prints","Reflective signage","Window branding","Vehicle branding / wraps","Fabric banners","Flags"],
"Signage & fabrication":["3D signs","Light boxes","Acrylic signs","Directional signs","Safety signs","Reception signs","Metal signs","Foam-board signs","Exhibition stands","Installation services"],
"Apparel & textiles":["DTF transfers","DTF application","Screen printing","Sublimation","Embroidery","Heat-transfer vinyl","T-shirts / polos","Hoodies","Jerseys","Workwear","Caps","Tote bags","Aprons","Uniforms"],
"Packaging & labels":["Product labels","Food labels","Cosmetic labels","Bottle labels","Paper bags","Gift bags","Boxes / cartons","Pouches","Sleeves","Wrapping paper","Tamper seals","Barcode / variable labels"],
"Promotional products":["Mugs","Tumblers","Water bottles","Pens","Keyholders","Lanyards","Wristbands","Notebooks","Umbrellas","Awards / plaques","Name tags","Phone accessories","Souvenir bundles"],
"Events & corporate branding":["Event identity","Stage / venue branding","Photo walls","Directional systems","Corporate stationery","Staff onboarding packs","Conference materials","Campaign materials","Recurring print-management service"],
"Materials & supplies for sale":["Blank apparel","DTF film","DTF powder","Ink / consumables","Vinyl rolls","Banner material","Paper / card stock","Packaging blanks","Other trade supplies"]}
for group,items in service_groups.items():
    subsection(group)
    for item in items: check(item+"   Offered: ☐ Now  ☐ Planned  ☐ Subcontracted")
field("Missing services","Add anything not listed above",3)

page_break(); section("4. Service master-data worksheet")
p("Duplicate this section for every service. The answers become the system’s service configurator, customer questions, production job card and pricing engine.")
for label,hint in [("Service name","Customer-facing name"),("Internal code / SKU","Unique identifier"),("Category and subcategory","Where customers find it"),("Plain-language description","What the customer receives"),("Typical use cases","Examples customers recognise"),("Target customers","Retail, corporate, reseller or all"),("Fulfilment type","Made in-house / subcontracted / hybrid"),("Responsible department","Production team or machine group"),("Standard turnaround","Business hours/days after artwork approval"),("Rush options","Cut-offs, surcharge and realistic lead time")]: field(label,hint)
subsection("Customer configuration questions")
table(["Question shown to customer","Input type","Allowed values / range","Required?","Affects price?"],[["Example: Finished width","Number + unit","100–5,000 mm","Yes","Yes"],["","","","☐","☐"],["","","","☐","☐"],["","","","☐","☐"],["","","","☐","☐"],["","","","☐","☐"]],[3000,1500,2500,1140,1800],8.2)
subsection("Options, dependencies and exclusions")
field("Available options","Size, colour, material, thickness/GSM, sides, finish, mounting, installation, etc.",3)
field("Option dependencies","Example: lamination only available for selected materials",2)
field("Invalid combinations","Options the system must prevent",2)
field("Maximum/minimum machine sizes","Printable area, margins and panel limits",2)
field("Customer guidance","Tooltips, photos, diagrams, warnings and sample uses",2)
subsection("Required digital assets")
for x in ["Hero / category image","At least three finished-work photographs","Material/colour swatches","Size diagram or measurement guide","Finishing-option examples","Before/after example","Short explanatory video (if useful)","Artwork template / dieline"]: check(x)

page_break(); section("5. Intelligent pricing-engine discovery")
callout("Pricing principle", "The customer price should be built from auditable components: design/prepress + setup + material + machine/production + finishing + labour + wastage + subcontracting + delivery + tax + margin − approved discount.")
subsection("Choose the calculation model for each service")
table(["Model","Typical use","Used for which services?"],[["Fixed price","Standard item or package with no meaningful variation",""],["Unit × quantity","Mugs, shirts, pens, labels",""],["Area-based","Banners, vinyl, wall graphics: width × height",""],["Linear length","Roll media, hemming, rails, sewing",""],["Sheet imposition","Cards/flyers: items per sheet + sheets required",""],["Machine-time","Engraving, embroidery, specialist production",""],["Stitch-count","Embroidery by estimated/actual stitches",""],["Colour/screen count","Screen printing and separations",""],["Tier / matrix","Quantity × size/material price grid",""],["Formula + review","Complex price calculated then staff-approved",""],["Quote only","No reliable instant price",""],["Hybrid","Base package plus configurable additions",""]],[1850,3900,4190],8.3)

subsection("Universal pricing inputs")
for x in ["Base/minimum order charge","Setup charge per job, design, colour, screen, plate, frame or machine","Material cost per sheet, roll, metre, square metre, piece, kilogram or set","Printable yield / imposition per sheet","Machine rate per impression, pass, minute or hour","Ink/coverage charge and white-ink surcharge","Labour rate per minute/hour or fixed operation","Wastage percentage and minimum spoilage quantity","Finishing charges: cutting, trimming, folding, binding, eyelets, hemming, lamination, mounting, sewing, packaging","Quantity discounts / price breaks","Rush surcharge","Small-order surcharge","Artwork/design charge","Delivery, installation and travel","Subcontractor cost and markup","VAT, levies or other taxes","Rounding rule and minimum payable amount","Target gross-margin floor","Maximum automatic discount"]: check(x)

subsection("Core formulas to confirm")
table(["Calculation","Proposed system formula","Confirm / amend"],[["Area","Width × height × number of panels × quantity (convert units)",""],["Material consumption","Required area or sheets ÷ usable yield, rounded up + wastage",""],["Imposition","Ceiling(quantity ÷ items per sheet) + spoilage sheets",""],["Print runs","Sheets × sides × passes/colours",""],["Embroidery","Digitising fee + garment + stitch-band price + extras × quantity",""],["DTF","Gang-sheet area/length + white ink/coverage + film + powder + press labour",""],["Selling price","Total direct cost ÷ (1 − target gross-margin %) + taxes",""],["Deposit","Eligible subtotal × deposit %; balance before dispatch",""],["Rush","Applicable production subtotal × rush % or fixed surcharge",""],["Delivery","Zone base fee + distance/weight/size increments",""]],[1800,5200,2940],8.4)

page_break(); section("6. Detailed price-component worksheet")
p("Complete one copy for every service. Every amount must state whether tax is included and when it becomes effective.")
field("Service / variant")
table(["Cost or price component","Unit","Supplier/direct cost","Customer price / rate","Minimum","Effective date"],[["Base/setup fee","per job","","","",""],["Material","","","","",""],["Printing/machine","","","","",""],["Ink / coverage","","","","",""],["Labour","","","","",""],["Finishing","","","","",""],["Packaging","","","","",""],["Installation","","","","",""],["Subcontracting","","","","",""],["Other","","","","",""]],[2600,1150,1650,1850,1200,1490],7.8)
subsection("Quantity breaks")
table(["From quantity","To quantity","Unit price / multiplier","Setup impact","Notes"],[["1","","","",""],["","","","",""],["","","","",""],["","","","",""],["","","","",""]],[1700,1700,2300,1800,2440],8.2)
subsection("Margin and approval controls")
field("Target gross margin %")
field("Absolute minimum gross margin %","Below this requires named approval")
field("Who can change prices?","Roles and approval workflow")
field("Who can discount and by how much?","Limits per role/customer type")
field("Price validity period","How long saved quotes remain valid")
field("Automatic review triggers","High value, low margin, unusual size, unavailable material, complex artwork",2)

page_break(); section("7. Material, machine and production capability master data")
subsection("Materials")
table(["Material","Unit","Supplier","Last cost","Usable dimensions","Wastage %","Reorder point"],[["","","","","","",""],["","","","","","",""],["","","","","","",""],["","","","","","",""],["","","","","","",""]],[1750,850,1450,1100,2000,1250,1540],7.7)
field("Material substitutions","Allowed alternatives and who approves them",2)
field("Colour / batch variation rules","What customers must be told",2)
subsection("Machines and work centres")
table(["Machine/work centre","Capabilities","Max size","Speed/capacity","Hourly cost","Operator roles","Downtime rules"],[["","","","","","",""],["","","","","","",""],["","","","","","",""],["","","","","","",""]],[1600,1900,1050,1500,1150,1400,1340],7.5)
field("Routing rules","Which machine is preferred by material, quantity, size and deadline",3)
field("Maintenance schedules","Planned downtime that affects promised dates",2)
field("Capacity constraints","Daily/weekly capacity by department and peak season",2)

section("8. Artwork, design and proofing rules")
for label,hint in [("Accepted file formats","PDF, AI, EPS, SVG, PSD, CDR, PNG, JPG, etc."),("Maximum upload size","Per file and total per order"),("Colour spaces","CMYK/RGB/Pantone rules"),("Resolution","Minimum DPI by output type"),("Bleed and safe area","By product type"),("Fonts and linked assets","Outline/embed/package requirements"),("Transparency / overprint rules","Preflight requirements"),("Variable data","CSV format, barcode and personalisation rules"),("Storage period","How long artwork is retained"),("Customer rights confirmation","Copyright, trademark and permission declaration")]: field(label,hint)
subsection("Artwork workflow")
for x in ["System performs automated file-type and size checks","Customer selects artwork-ready, needs adjustment, or needs full design","Staff performs preflight and records issues","Customer receives digital proof","Customer approval is timestamped and versioned","Production cannot start before approval unless explicitly authorised","Any change after approval creates a new version and may affect price/deadline","Final production files are access-controlled"]: check(x)
field("Proof approval wording","Exact legal/customer confirmation",2)
field("Included revisions","Number, scope and turnaround")
field("Additional revision pricing","Rate and approval process")

page_break(); section("9. Customer ordering experience")
subsection("Customer account and checkout")
for x in ["Guest checkout","Customer account","Business/corporate account","Saved addresses","Saved artwork library","Reorder previous job","Save configuration as draft","Request formal quote","Instant price and pay","Purchase-order / credit-account ordering","Tax invoice and receipt download","Order status timeline","In-platform messages","WhatsApp/SMS/email notifications"]: check(x)
field("Required customer fields","Individual and company details",2)
field("Customer verification","Email, phone OTP, business verification or none")
field("Age/content restrictions","Restricted designs or products")
subsection("Order configuration flow")
table(["Step","Customer action","System response","Human review?"],[["1","Choose service or shop product","Show use cases, price basis and turnaround",""],["2","Enter specifications","Validate dimensions/options and update price",""],["3","Upload artwork / request design","Preflight or add design service",""],["4","Choose proof, delivery and deadline","Calculate fulfilment promise and fees",""],["5","Review price breakdown","Show tax, discount, deposit/balance",""],["6","Pay securely","Initialize Paystack transaction",""],["7","Confirmation","Create order/job card and notify team",""],["8","Track and approve","Status, proof approval, balance payment",""],["9","Receive and confirm","Delivery/pickup confirmation and feedback",""]],[650,2800,4600,1890],8.1)

section("10. Online shop / stock-product requirements")
field("Product families","DTF films, blank/branded apparel, ready-made items, consumables, etc.",2)
table(["Product data","Required details"],[["Identity","Name, SKU, barcode, brand, category, description"],["Variants","Size, colour, material, pack size and variant SKU"],["Price","Cost, selling price, sale price, tax and wholesale tiers"],["Inventory","Stock on hand, reserved, available, reorder level and location"],["Media","Primary image, gallery, video, swatch images"],["Fulfilment","Weight, dimensions, pickup/delivery availability and lead time"],["Policies","Returns eligibility, warranty and usage/safety information"],["Related items","Upsells, bundles, consumables and compatible products"]],[2300,7640],8.7)
field("Inventory reservation rule","When stock becomes reserved and when reservation expires")
field("Overselling / backorder policy")
field("Wholesale and reseller pricing rules",lines=2)

page_break(); section("11. Paystack, Mobile Money and payment operations")
callout("Security boundary", "Vikipat must never collect or store card or Mobile Money credentials. Payment details remain on Paystack-hosted/approved interfaces. The platform stores references, status, amounts and audit records only.")
for label,hint in [("Paystack business account owner","Legal entity and responsible person"),("Paystack public key","Provide securely during implementation—not in this workbook"),("Paystack secret key","Provide through a secrets manager—not email or this workbook"),("Settlement bank account","Account holder and finance owner; avoid full bank details here"),("Settlement frequency","Expected Paystack settlement cycle"),("Enabled channels","Mobile Money, cards, bank, USSD, QR or others"),("Accepted currency","GHS and any future currencies"),("Transaction fee treatment","Absorbed, passed on where lawful, or included in pricing")]: field(label,hint)
subsection("Payment rules")
for x in ["Full payment before production","Percentage deposit, balance before production","Percentage deposit, balance before dispatch","Staff-approved pay-later corporate account","Split payments / payment links","Additional payment after scope change","Refund to original channel","Store credit","Cash/manual bank payment recorded by finance"]: check(x)
field("Deposit rules","Services, percentages, thresholds and exceptions",2)
field("Balance-payment triggers","Proof approval, production completion, dispatch, etc.",2)
field("Failed/abandoned payment handling","Retry window and stock/order reservation",2)
field("Webhook reconciliation rules","How duplicate, delayed or mismatched events are handled",2)
field("Refund approval matrix","Who may approve by amount/reason",2)
field("Chargeback/dispute process","Evidence, owner and customer communication",2)
subsection("Payment status model")
table(["Status","Meaning","Allowed next action"],[["Pending","Transaction initialized, not confirmed","Retry or expire"],["Paid","Paystack-verified amount/reference","Release order to workflow"],["Part-paid","Valid deposit received","Collect balance at configured gate"],["Failed","Provider reports failure","Retry with new reference"],["Expired","Payment window ended","Release reservation"],["Refund pending","Approved and submitted","Track provider outcome"],["Refunded","Funds returned","Close financial adjustment"],["Disputed","Chargeback/dispute opened","Freeze and investigate"]],[1900,3900,4140],8.4)

page_break(); section("12. Back-office staff operations")
subsection("Staff roles and permissions")
table(["Role","Core responsibilities","May view","May change / approve"],[["Customer service / sales","Validate orders, communicate, quote exceptions","Customer/order data","Specifications within limits"],["Designer / prepress","Prepare artwork, proofs and production files","Assigned artwork/jobs","Artwork versions and proof status"],["Production supervisor","Plan, assign and monitor work","All production jobs/capacity","Routing, priority, due date within policy"],["Machine operator","Execute assigned tasks and report consumption","Assigned job card","Task status, actual usage, issue logs"],["Quality control","Inspect output and record pass/fail","Job and QC criteria","QC decision, rework request"],["Storekeeper","Issue materials and manage stock","Inventory and reservations","Receipts, issues and adjustments"],["Dispatch","Prepare pickup/delivery and capture proof","Ready orders/contact/address","Dispatch and delivery status"],["Finance","Reconcile payments, refunds and expenses","Financial data","Refund workflow and reconciliation"],["System administrator","Users, roles and configuration","All permitted records","Configuration; no silent audit deletion"],["Executive admin","Business-wide oversight","All dashboards/reports","Approvals and policy controls"]],[1750,2850,2200,3140],7.8)

subsection("Production order lifecycle")
statuses=[("Received","Payment verified; order record locked to paid specification"),("Review required","Staff checks artwork, feasibility, price exception or stock"),("Awaiting customer","Need clarification, corrected artwork, proof approval or balance"),("Approved","Ready for production planning"),("Scheduled","Assigned department/machine/operator and time slot"),("In production","Work started; actual time/material consumption recorded"),("Quality check","Output inspected against configured checklist"),("Rework","Issue identified; reason, responsibility and cost recorded"),("Ready","Packed and ready for pickup/dispatch"),("Out for delivery","Assigned to rider/courier with tracking"),("Delivered / collected","Proof of delivery or pickup captured"),("Closed","Financial and operational checks complete"),("Cancelled","Reason, approval and refund outcome recorded")]
table(["Status","Definition / required system behaviour"],statuses,[2300,7640],8.4)
field("Status changes customers may see","Customer-friendly wording")
field("Internal statuses customers must not see","Fraud, margin, supplier or staff notes")
field("Escalation rules","Late jobs, machine breakdown, artwork delay, failed QC, stockout",3)

page_break(); section("13. Job cards, scheduling and production capture")
subsection("Digital job card fields")
for x in ["Order/job number and barcode/QR code","Customer and contact details","Paid specification snapshot","Artwork version and proof approval timestamp","Material, colour, dimensions, quantity and finishing","Department, machine and assigned staff","Planned start, due time and priority","Step-by-step work instructions","Expected material and labour","Actual material used, waste, time and output","QC checklist and evidence/photos","Rework reason and approval","Packing, pickup/delivery instructions","Internal notes and customer-visible notes"]: check(x)
subsection("Scheduling decisions")
field("How jobs are prioritised","Paid time, due date, rush, VIP, machine batching, material",2)
field("Batching rules","Group by material, colour, machine, ink or finishing",2)
field("Staff shifts and capacity","People, hours, breaks and overtime",2)
field("Dependencies","Artwork → print → cure/dry → finish → QC → dispatch",2)
field("Late-order alert thresholds","When and whom to alert")
field("Production board views","Department, machine, operator, day/week, due date, status",2)

section("14. Quality control, rework and waste")
table(["Service / category","QC checks","Tolerance","Evidence required","Approver"],[["Large format","Size, colour, finish, eyelets, damage","","Photo / measurement",""],["Small format","Quantity, trim, colour, pagination, finish","","Sample / count",""],["Apparel / DTF","Placement, adhesion, colour, garment defect","","Photo / wash test",""],["Embroidery","Placement, thread colour, stitches, puckering","","Photo / sample",""],["Packaging / labels","Dimensions, adhesion, cut, text/barcode","","Photo / scan",""],["Other","","","",""]],[1750,3000,1300,2200,1690],7.8)
field("Reprint/rework policy","When Vikipat bears cost versus customer change",2)
field("Waste reasons","Setup, spoilage, machine fault, operator error, material defect, customer change",2)
field("Approval thresholds","Who authorises rework by value/quantity")
field("How recovered/scrap material is handled",lines=2)

page_break(); section("15. Delivery, installation and pickup")
for x in ["Workshop pickup","Vikipat delivery team","Third-party courier","Customer-arranged transport","Nationwide delivery","International delivery","On-site installation"]: check(x)
table(["Zone / method","Base fee","Distance/weight/size rule","Lead time","Free-delivery threshold"],[["Mallam–Gbawe / nearby","","","",""],["Accra zone 1","","","",""],["Accra zone 2","","","",""],["Outside Accra","","","",""],["Pickup","","","",""]],[1900,1300,3100,1600,2040],8)
field("Address fields","Digital address, landmark, GPS pin, recipient and phone")
field("Bulky/fragile restrictions","Vehicle, packaging and handling rules",2)
field("Installation pricing","Call-out, travel, labour, height/equipment, permits",2)
field("Proof of delivery","OTP, signature, photo, name and timestamp")
field("Failed delivery / storage fees","Retries, customer unavailable and uncollected work",2)

section("16. Customer communication and notifications")
table(["Event","Email","WhatsApp/SMS","In-app","Internal recipients"],[["Account created","☐","☐","☐",""],["Order/payment confirmed","☐","☐","☐",""],["Artwork problem","☐","☐","☐",""],["Proof ready / approved","☐","☐","☐",""],["Balance due","☐","☐","☐",""],["Production started","☐","☐","☐",""],["Delay / issue","☐","☐","☐",""],["Ready for pickup","☐","☐","☐",""],["Out for delivery","☐","☐","☐",""],["Delivered","☐","☐","☐",""],["Refund update","☐","☐","☐",""],["Review / reorder reminder","☐","☐","☐",""]],[2900,1200,1750,1200,2890],8)
field("Message tone and sender name")
field("WhatsApp provider / API","Official business account details and template approval status")
field("SMS/email provider","Existing account or preferred provider")
field("Escalation contacts","Operational, payment and security incidents",2)

page_break(); section("17. Financial intelligence and executive oversight")
callout("Financial truth", "Revenue must be recognised from verified payment/order events, while refunds, discounts, taxes, payment fees, material consumption, waste and rework remain separately visible. No dashboard should treat gross collections as profit.")
subsection("Executive dashboard")
for x in ["Gross sales, net sales and collections","Orders, average order value and conversion","Deposits collected and outstanding balances","Revenue by service, category, channel and customer type","Direct cost, gross profit and gross margin","Discounts by staff/customer/reason","Paystack fees and settlement reconciliation","Tax collected and tax liability","Refunds, cancellations, disputes and chargebacks","Work in progress value","Inventory value and stock variance","Waste/rework cost and causes","Machine utilisation and staff throughput","On-time delivery and production cycle time","Accounts receivable / credit customers","Top customers and repeat-purchase rate","Forecast cash inflows and production demand"]: check(x)
subsection("Accounting and controls")
field("Accounting basis","Cash/accrual and revenue-recognition policy")
field("Chart of accounts / categories","Revenue, COGS, expenses, taxes, liabilities",2)
field("Tax rules","VAT registration, rates, inclusive/exclusive pricing and invoice fields",2)
field("Paystack settlement reconciliation","Frequency, owner and exception handling",2)
field("Expense capture","Purchases, petty cash, subcontractors, delivery, maintenance",2)
field("Purchase-order workflow","Request, approval, supplier order, receipt and invoice match",2)
field("Corporate credit terms","Eligibility, limits, ageing, statements and collections",2)
field("Period close","Daily cash-up, weekly review, month-end lock and adjustments",2)
table(["Approval","Thresholds","Primary approver","Backup","Evidence"],[["Discount","","","",""],["Refund","","","",""],["Price override","","","",""],["Stock adjustment","","","",""],["Purchase order","","","",""],["Expense","","","",""],["Credit limit","","","",""]],[1900,2100,2000,1700,2240],8)

page_break(); section("18. Reporting and data exports")
table(["Report","Frequency","Filters / dimensions","Recipients","Export"],[["Daily orders and production","Daily","Status, department, due date","","PDF/Excel"],["Sales and margin","Daily/weekly/monthly","Service, channel, customer","","Excel"],["Payments and settlements","Daily","Provider status/reference","","Excel/CSV"],["Inventory movement","Weekly","SKU, location, reason","","Excel"],["Waste and rework","Weekly","Service, machine, staff, reason","","PDF/Excel"],["Tax and invoices","Monthly","Tax period/status","","Excel/CSV"],["Executive performance","Monthly","Current vs target/prior period","","PDF"],["Customer retention","Monthly","Cohort, repeat, frequency","","Excel"]],[2200,1200,2800,1900,1840],7.8)
field("Existing accounting software","Name, edition and integration/export needs")
field("Existing spreadsheets/reports to reproduce","Attach examples")
field("Who may see cost, margin and profit data?","Specify role-level restrictions",2)

section("19. Users, access control and security")
for x in ["Unique user account for every employee","Role-based least-privilege access","Multi-factor authentication for administrators and finance","Password and session policy","Immediate user deactivation on exit","Device/session visibility and logout","Immutable audit log for sensitive actions","Approval workflow for high-risk changes","Encryption in transit and at rest","Daily encrypted backups and tested restoration","Malware scanning and file-type validation for uploads","Rate limiting, bot protection and abuse monitoring","Security incident contacts and response procedure","Data-retention and deletion policy","Customer consent and privacy notice","Periodic permission and audit review"]: check(x)
field("Staff list","Name, role, department, email/phone and required access",3)
field("Sensitive data restrictions","Costs, margins, customer lists, artwork, bank/settlement and employee data",2)
field("Backup recovery targets","Maximum acceptable data loss and downtime")
field("Legal/privacy requirements","Ghana Data Protection Act obligations, terms and policies to confirm with counsel",2)

page_break(); section("20. Integrations and external services")
table(["Capability","Existing/preferred provider","Account owner","Required data flow","Phase"],[["Payments","Paystack","","Transactions, webhooks, refunds, settlements","1"],["WhatsApp","","","Order/customer notifications",""],["SMS","","","Transactional notifications",""],["Email","","","Receipts, proofs, reports",""],["Maps/delivery","","","Address, distance and tracking",""],["Accounting","","","Invoices, payments, expenses, tax",""],["Cloud storage","","","Artwork and proof storage",""],["Analytics","","","Conversion and product analytics",""],["Courier","","","Rates, booking and tracking",""],["Identity/OTP","","","Phone/email verification",""],["Barcode/printing","","","Job cards, labels and receipts",""]],[1700,2150,1650,3000,1440],7.8)
field("Existing domains and hosting","Registrar, DNS owner and access process")
field("Existing social/store channels","Instagram, Facebook, Google Business, marketplace, etc.")
field("Legacy data to import","Customers, products, prices, artwork, orders, inventory",2)

section("21. Content, brand and customer experience assets")
for x in ["Primary and alternate logos in vector format","Brand-colour specifications (HEX/CMYK/Pantone)","Approved typefaces and licences","Brand guidelines","Service photographs—original, high resolution","Team/workshop/machine photographs","Customer testimonials and permission","Case studies and finished-project details","Company story and leadership bio","Terms and conditions","Privacy and cookie policy","Refund/cancellation policy","Delivery/pickup policy","Artwork and colour disclaimer","FAQ and customer-support scripts","Social-media links and handles"]: check(x)
field("Languages","English and any planned local-language support")
field("Accessibility requirements","Known customer needs and compliance target")
field("Competitors / inspiration","What to emulate and what to avoid",2)

page_break(); section("22. Exception scenarios and business rules")
p("Explain the required response, owner, customer message and financial outcome for each scenario.")
scenarios=["Customer pays wrong amount","Paystack says paid but webhook is delayed","Duplicate payment","Customer changes quantity/specification after payment","Uploaded artwork is unusable","Customer does not approve proof","Material becomes unavailable","Machine breaks after order acceptance","Job will miss promised deadline","Output fails QC","Customer rejects colour despite approved proof","Partial production completed before cancellation","Customer requests refund","Delivered item is damaged","Courier loses order","Customer does not collect finished work","Corporate customer exceeds credit limit","Suspicious/fraudulent order","Staff enters incorrect price or discount","System is unavailable during business hours"]
table(["Scenario","Required action / owner / customer outcome"],[[s,""] for s in scenarios],[3600,6340],8)

section("23. Sample-job evidence pack")
p("Provide at least five completed jobs for every major service family. Remove sensitive customer details if necessary, but retain all pricing and production figures.")
table(["Evidence item","Attach / complete"],[["Original customer request","☐"],["Final specification and quantity","☐"],["Artwork / proof versions","☐"],["Original quote and final invoice","☐"],["Detailed cost calculation","☐"],["Materials consumed and waste","☐"],["Machine and labour time","☐"],["Payment method and fees","☐"],["Production timeline and status changes","☐"],["Delivery cost / proof of delivery","☐"],["Discount, rework or exception","☐"],["Actual gross margin","☐"]],[4700,5240],8.5)

page_break(); section("24. Data-preparation templates")
subsection("Service import columns")
p("Service code • service name • category • description • calculation model • minimum charge • setup fee • unit • turnaround • rush rule • tax class • active status • review threshold • responsible department")
subsection("Product/variant import columns")
p("SKU • barcode • product • variant • category • cost price • selling price • tax • stock • reorder level • weight • dimensions • supplier • image filenames • active status")
subsection("Material import columns")
p("Material code • name • supplier • unit • unit cost • roll/sheet dimensions • usable yield • wastage • compatible machines • reorder level • lead time")
subsection("Price-break import columns")
p("Service/SKU • option combination • minimum quantity • maximum quantity • rate • setup adjustment • effective date • expiry date • customer tier")
subsection("Customer import columns")
p("Customer ID • type • legal/trading name • contact • phone • email • billing/delivery addresses • tax ID • credit status/limit • price tier • outstanding balance • consent fields")
callout("File format", "Provide clean Excel or CSV files with one record per row, stable IDs, separate columns for separate facts, no merged cells, and a data dictionary explaining codes and units.")

section("25. Prioritisation and phased launch")
table(["Capability","Must launch","Phase 2","Later","Notes / dependency"],[["Service configurator and intelligent pricing","☐","☐","☐",""],["Paystack / Mobile Money","☐","☐","☐",""],["Customer accounts and reorders","☐","☐","☐",""],["Online stock-product shop","☐","☐","☐",""],["Artwork and proof approval","☐","☐","☐",""],["Production job board","☐","☐","☐",""],["Inventory and purchasing","☐","☐","☐",""],["QC, waste and rework","☐","☐","☐",""],["Delivery and tracking","☐","☐","☐",""],["Financial intelligence","☐","☐","☐",""],["Corporate accounts / credit","☐","☐","☐",""],["WhatsApp automation","☐","☐","☐",""],["Accounting integration","☐","☐","☐",""],["Advanced capacity scheduling","☐","☐","☐",""]],[3400,1150,1150,1000,3240],8)
field("Desired launch date","Include immovable events/deadlines")
field("Available internal team","People assigned to discovery, testing, content and migration",2)
field("Budget range and payment milestones","Optional but strongly recommended",2)

page_break(); section("26. Definition of done and acceptance")
for x in ["Every launch service has an approved configuration and pricing formula","At least 20 real historic jobs reproduce the expected selling price within agreed tolerance","All payment statuses, webhook retries and duplicate events are tested","No order enters production without the configured payment/proof gates","Role permissions are tested with real staff scenarios","Production staff can complete job cards on their actual devices","Finance can reconcile orders, Paystack transactions, settlements, refunds and fees","Inventory reservation, consumption, waste and adjustment paths are tested","Customer notifications use approved wording and sender accounts","Backup restoration and security incident procedures are tested","Legal policies and tax/invoice requirements are approved by qualified advisers","Management approves dashboards, margin definitions and period-close reports","A pilot group completes end-to-end orders before public launch"]: check(x)
field("Pricing accuracy tolerance","Example: exact or within stated rounding rule")
field("Performance expectations","Peak concurrent users, response time and upload sizes")
field("Pilot customers and staff","Names/roles")
field("Final acceptance authority","Name and role")

section("27. Open decisions and assumptions log")
table(["ID","Decision / assumption","Owner","Due date","Status / outcome"],[[str(i),"","","","Open"] for i in range(1,13)],[700,4300,1700,1400,1840],8)

section("28. Final sign-off")
p("By signing, reviewers confirm that the information supplied is sufficiently complete for solution design and that omissions or later changes may affect price, scope and delivery dates.")
table(["Role","Name","Signature","Date"],[["Managing director / owner","","",""],["Operations lead","","",""],["Finance lead","","",""],["Project lead","","",""]],[2600,2500,2900,1940],8.5)
callout("Secure handover reminder", "Do not place live Paystack secret keys, bank credentials, passwords or government-issued personal identifiers in this workbook. Transfer secrets only through an agreed secure channel during implementation.", "FFF4D8")

# Keep table rows readable across pages and repeat header rows.
for t in doc.tables:
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row in t.rows:
        trPr=row._tr.get_or_add_trPr(); cant=OxmlElement("w:cantSplit"); trPr.append(cant)

# Metadata
doc.core_properties.title="Vikipat Digital Commerce & Production Operations Platform — Discovery Workbook"
doc.core_properties.subject="Service catalogue, intelligent pricing, Paystack, production, fulfilment and financial requirements"
doc.core_properties.author="Vikipat Platform Project"
doc.core_properties.keywords="Vikipat, printing, branding, pricing, production, Paystack, ecommerce, requirements"
doc.save(OUT)
print(OUT)
