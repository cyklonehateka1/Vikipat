from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path(__file__).resolve().parents[1]/"docs"/"Vikipat_Core_Online_Order_and_Pricing_Discovery_Workbook.docx"
OUT.parent.mkdir(parents=True,exist_ok=True)
NAVY="08194A"; BLUE="1234D8"; INK="17213A"; MUTED="5A657A"; PALE="EDF3FF"; LIGHT="F6F8FC"; WHITE="FFFFFF"; BORDER="CCD5E3"; GOLD="FFF3D4"
d=Document(); s=d.sections[0]; s.page_width=Inches(8.5); s.page_height=Inches(11); s.top_margin=Inches(.72); s.bottom_margin=Inches(.68); s.left_margin=Inches(.78); s.right_margin=Inches(.78); s.header_distance=Inches(.3); s.footer_distance=Inches(.32)

def fmt(r,size=10.5,bold=False,color=INK,italic=False,name="Aptos"):
    r.font.name=name; r._element.get_or_add_rPr().rFonts.set(qn("w:ascii"),name); r._element.rPr.rFonts.set(qn("w:hAnsi"),name); r.font.size=Pt(size); r.bold=bold; r.italic=italic; r.font.color.rgb=RGBColor.from_string(color)

n=d.styles["Normal"]; n.font.name="Aptos"; n.font.size=Pt(10.5); n.font.color.rgb=RGBColor.from_string(INK); n.paragraph_format.space_after=Pt(6); n.paragraph_format.line_spacing=1.16
for k,z,c,b,a in [("Heading 1",18,NAVY,18,8),("Heading 2",13.5,BLUE,13,6),("Heading 3",11.5,NAVY,9,4)]:
    x=d.styles[k]; x.font.name="Aptos Display"; x.font.size=Pt(z); x.font.bold=True; x.font.color.rgb=RGBColor.from_string(c); x.paragraph_format.space_before=Pt(b); x.paragraph_format.space_after=Pt(a); x.paragraph_format.keep_with_next=True

def shade(c,fill):
    pr=c._tc.get_or_add_tcPr(); x=pr.find(qn("w:shd"));
    if x is None: x=OxmlElement("w:shd"); pr.append(x)
    x.set(qn("w:fill"),fill)
def width(c,w):
    pr=c._tc.get_or_add_tcPr(); x=pr.find(qn("w:tcW"));
    if x is None: x=OxmlElement("w:tcW"); pr.append(x)
    x.set(qn("w:w"),str(w)); x.set(qn("w:type"),"dxa")
def margin(c):
    pr=c._tc.get_or_add_tcPr(); m=pr.find(qn("w:tcMar"));
    if m is None: m=OxmlElement("w:tcMar"); pr.append(m)
    for side,val in [("top",110),("start",130),("bottom",110),("end",130)]:
        x=m.find(qn("w:"+side));
        if x is None: x=OxmlElement("w:"+side); m.append(x)
        x.set(qn("w:w"),str(val)); x.set(qn("w:type"),"dxa")
def para(text="",size=10.5,bold=False,color=INK,italic=False,align=None,after=6):
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(after)
    if align is not None:p.alignment=align
    fmt(p.add_run(text),size,bold,color,italic); return p
def field(label,hint="",lines=1):
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(2); fmt(p.add_run(label+": "),10.3,True,NAVY)
    if hint:fmt(p.add_run(hint),9.3,False,MUTED,True)
    for _ in range(lines):para("________________________________________________________________________________",8,False,BORDER,after=3)
def check(text):
    p=d.add_paragraph(); p.paragraph_format.left_indent=Inches(.2); p.paragraph_format.first_line_indent=Inches(-.2); p.paragraph_format.space_after=Pt(3); fmt(p.add_run("☐ "),11,True,BLUE); fmt(p.add_run(text),10.1)
def table(headers,rows,widths=None,fs=8.2):
    t=d.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; widths=widths or [9940//len(headers)]*len(headers)
    tw=t._tbl.tblPr.find(qn("w:tblW")); tw.set(qn("w:w"),"9940"); tw.set(qn("w:type"),"dxa"); ind=OxmlElement("w:tblInd"); ind.set(qn("w:w"),"120"); ind.set(qn("w:type"),"dxa"); t._tbl.tblPr.append(ind)
    for i,c in enumerate(t.rows[0].cells): shade(c,NAVY); width(c,widths[i]); margin(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); fmt(p.add_run(headers[i]),fs,True,WHITE)
    for ri,row in enumerate(rows):
        cs=t.add_row().cells
        for i,v in enumerate(row):
            c=cs[i]; width(c,widths[i]); margin(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2:shade(c,LIGHT)
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); fmt(p.add_run(str(v)),fs)
    h=OxmlElement("w:tblHeader"); t.rows[0]._tr.get_or_add_trPr().append(h); d.add_paragraph().paragraph_format.space_after=Pt(1); return t
def callout(title,text,fill=PALE):
    t=d.add_table(rows=1,cols=1); t.autofit=False; c=t.cell(0,0); width(c,9940); margin(c); shade(c,fill); p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(3); fmt(p.add_run(title),10.5,True,NAVY); p=c.add_paragraph(); p.paragraph_format.space_after=Pt(0); fmt(p.add_run(text),9.8); t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader")); d.add_paragraph().paragraph_format.space_after=Pt(1)
def h1(x,lead=""): d.add_heading(x,1); para(lead,9.8,False,MUTED) if lead else None
def h2(x): d.add_heading(x,2)
def pb():d.add_page_break()

hp=s.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; fmt(hp.add_run("VIKIPAT / ONLINE ORDER DISCOVERY"),8,True,MUTED)
fp=s.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; fmt(fp.add_run("Confidential • Vikipat Media Solutions • "),8,False,MUTED); fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),"PAGE"); fp._p.append(fld)

para("VIKIPAT",13,True,BLUE,align=WD_ALIGN_PARAGRAPH.CENTER,after=50)
para("Online Ordering &\nPricing Discovery Workbook",30,True,NAVY,align=WD_ALIGN_PARAGRAPH.CENTER,after=12)
para("The practical information needed to quote, accept and fulfil customer storefront orders correctly",14,False,BLUE,align=WD_ALIGN_PARAGRAPH.CENTER,after=30)
callout("Scope", "One Vikipat branch. Online customer orders only. The existing walk-in system remains outside this project unless a small hand-off is later requested.")
para("This is not a general company questionnaire. It concentrates on the core business knowledge the application needs: what Vikipat sells, what customers must specify, how each price is calculated, when staff review is required, and what happens when an order does not follow the happy path.",10.5,False,INK,align=WD_ALIGN_PARAGRAPH.CENTER,after=32)
table(["Completed by","Operations checked","Pricing checked","Approved"],[["","","",""]],[2485]*4,8.5)
pb()

h1("1. What will be available for online ordering?","Tick only services Vikipat is ready to accept and fulfil from the storefront. Add missing services; do not document walk-in-only work.")
groups={"Large format":["Banners","Pull-up banners","Backdrops/photo walls","Vinyl/window branding","Posters","Vehicle branding","Signage/boards","Installation"],"Small format":["Business cards","Flyers","Brochures","Letterheads/envelopes","Books/programmes/manuals","Invitations/cards","Stickers and labels","Receipt/invoice books"],"Apparel and branding":["DTF transfers only","DTF printed apparel","Screen printing","Embroidery","Sublimation","Caps","Uniforms/workwear","Tote bags"],"Packaging and promotional items":["Product labels","Paper/gift bags","Boxes/packaging","Mugs/tumblers/bottles","Pens/keyholders/lanyards","Awards/plaques","Event/corporate packs"],"Ready-made goods and supplies":["DTF film/powder","Blank apparel","Ready-branded shirts/items","Vinyl/banner material","Other consumables"]}
for g,items in groups.items():
    h2(g)
    for x in items:check(x+"   ☐ Instant price   ☐ Staff-confirmed price   ☐ Not online")
field("Anything missing","List only customer-storefront services",2)
callout("Important decision", "Instant price means the system may safely accept payment without a person checking the job first. Use staff-confirmed price where artwork, site conditions, unusual materials or production complexity can change the true cost.",GOLD)

pb(); h1("2. Complete one service sheet for every online service","This is the most important part of the workbook. Copy these pages for each service whose options or calculation differ.")
h2("A. Service definition")
for a,b in [("Service name","Exact customer-facing name"),("What the customer receives","Simple description and what is included"),("Who normally orders it","Useful customer examples"),("Produced by Vikipat or outsourced?","If outsourced, state which part"),("Normal turnaround","When the clock starts: payment, artwork approval, etc."),("Rush turnaround","Whether offered and the extra charge"),("Minimum order","Minimum quantity, size or value")]:field(a,b)
h2("B. Questions the storefront must ask")
table(["Question / option","Example answers","Required?","Changes price?","Invalid combinations / limits"],[["Quantity","1, 10, 50, 100","Yes","Yes",""],["Finished size","Width × height + unit","","",""],["Material","PVC, vinyl, paper, cotton…","","",""],["Print sides / colours","One/two sides; full colour…","","",""],["Finishing","Lamination, eyelets, cutting…","","",""],["Artwork/design","Ready file, adjustment, full design","","",""],["Delivery/installation","Pickup, delivery, installation","","",""],["Other","","","",""],["Other","","","",""]],[2100,2250,950,1150,3490],7.7)
field("Customer guidance needed","Explain terms customers commonly misunderstand",2)
field("Options the system must block","Impossible sizes/materials/finishes or unsafe combinations",2)
field("When staff must review before payment","Exact triggers—not just 'complex jobs'",2)

pb(); h1("3. How is this service priced?","Write the real calculation used by the person who quotes jobs today. Separate every component so the system can explain and audit the estimate.")
field("Service name")
h2("A. Calculation method")
for x in ["Fixed price","Price per item × quantity","Width × height × rate per square unit","Price per linear metre","Sheets required based on how many fit per sheet","Machine time / production time","Embroidery stitch count","Number of colours/screens/passes","Quantity-and-option price table","Supplier cost + markup","Other formula"]:check(x)
field("Exact formula in plain language","Example: area × material rate + print rate + eyelets + design + delivery",3)
h2("B. Every component used in the price")
table(["Component","Charged per","Cost to Vikipat","Customer rate","Minimum / rounding rule"],[["Setup / minimum job","job","","",""],["Material","sheet/metre/m²/item","","",""],["Printing / machine","pass/sheet/minute","","",""],["Ink / colour / coverage","","","",""],["Labour","operation/hour","","",""],["Design / artwork correction","","","",""],["Finishing","","","",""],["Packaging","","","",""],["Outsourced work","","","",""],["Other","","","",""]],[2300,1750,1750,1800,2340],7.8)
field("Waste/spoilage allowance","Percentage, fixed pieces/sheets, or rule by quantity")
field("Quantity discounts","Breakpoints and exact rate/multiplier",2)
field("Urgent-order surcharge","Percentage or fixed amount and qualifying deadline")
field("Margin/markup rule","How selling price is derived from cost")
field("VAT/tax treatment","Included or added; confirm with accountant")
field("Final rounding","Nearest pesewa/cedi/5/10, always up, etc.")

pb(); h1("4. Give real price examples","For each major service, provide recent examples. These are essential for testing whether the pricing engine matches how Vikipat actually quotes.")
table(["Example","Customer specification","How each part was calculated","Final price","Was anything overridden?"],[["1","","","",""],["2","","","",""],["3","","","",""],["4","","","",""],["5","","","",""]],[800,2700,3400,1300,1740],7.8)
h2("For each example, attach")
for x in ["Customer request/specification","Artwork or representative file","Original quote/invoice","Material and finishing used","Calculation notes or price sheet","Any discount, rush charge, wastage or correction"]:check(x)
field("Acceptable calculation difference","Must match exactly, or define an allowed rounding tolerance")
callout("Why examples matter", "Written formulas often miss a rule the estimator applies from experience. Real jobs expose hidden minimums, yield decisions, waste, difficult artwork and judgement calls.")

h1("5. Artwork, design and proof rules","Only collect rules that affect whether the order can be priced, accepted or released to production.")
for a,b in [("Accepted upload formats","PDF, AI, CDR, EPS, PSD, PNG, JPG, etc."),("Maximum upload size","Per file and total"),("Minimum resolution","By service where different"),("Colour requirement","CMYK, Pantone, RGB warning"),("Bleed/safe area","By applicable service"),("Maximum/minimum printable size","Include machine margins"),("Artwork-check responsibility","What the system checks vs staff"),("Design charges","Minor adjustment, recreation and full design"),("Included revisions","Number and meaning of a revision"),("Proof required","Which services need customer approval before production")]:field(a,b)
for x in ["Customer uploads print-ready artwork","Customer needs Vikipat to correct artwork","Customer needs a new design","Staff rejects file and requests replacement","Staff sends proof; customer approves","Customer changes specification after approving proof"]:check(x+" — describe price/deadline effect if applicable")
field("Proof approval statement","Exact message customer must accept",2)

pb(); h1("6. Online order journey and production hand-off","Confirm the minimum workflow for orders coming from the storefront.")
table(["Stage","What must happen","Who acts","Customer sees"],[["Configure","Customer answers service questions; estimate updates","Customer/system","Price and turnaround"],["Review gate","System decides instant payment or staff confirmation","System/staff","Pay now or awaiting review"],["Payment","Paystack verifies payment/deposit","Customer/system","Payment confirmation"],["Artwork/proof","File checked, corrections/proof approved","Designer/customer","Action needed / approved"],["Accepted","Paid and production-ready order becomes a job","Operations","Order accepted"],["Production","Job assigned and worked on","Production staff","In production"],["Quality check","Finished work checked; rework if needed","QC/staff","Preparing order"],["Ready","Packed for pickup/delivery","Staff","Ready"],["Completed","Picked up/delivered with confirmation","Staff/customer","Completed"]],[1450,3900,1900,2690],8)
field("Staff dashboard queues needed","Example: needs review, awaiting artwork, ready to produce, due today",2)
field("Information staff need on the job card","Specification, artwork, payment, deadline, notes, etc.",3)
field("Internal status names","If different from the proposed flow")
field("Customer-friendly status names","Avoid exposing internal production details")
field("Who may change price/specification after payment?","And what approval is required")

pb(); h1("7. Edge cases: what should the system do?","For each scenario, state whether to block checkout, recalculate, request staff review, collect/refund money, change deadline, or cancel.")
cases=["Size is below/above machine or material limits","Quantity is below minimum","Selected options cannot be produced together","Required material/colour is out of stock","Price changes while customer has an unpaid basket","Customer enters the wrong unit (mm/cm/inches)","Artwork dimensions do not match ordered size","Artwork is low resolution, corrupt or wrong format","Customer has no artwork and needs design","Customer asks for changes after proof approval","Customer changes size/quantity/material after payment","Final reviewed price is higher than the estimate","Final reviewed price is lower than the estimate","Customer pays only part of required amount","Payment succeeds but confirmation/webhook is delayed","Customer pays twice","Payment fails after stock/capacity was reserved","Order is urgent but capacity is unavailable","Machine breaks or production will be late","More material/waste is needed than estimated","Finished work fails quality check","Part of the order is good and part must be redone","Customer cancels before production","Customer cancels after artwork/design has started","Customer cancels after production has started","Vikipat cannot fulfil the order","Customer does not approve proof/respond","Customer does not collect finished work","Delivery address/cost changes after payment","Delivery fails or item is damaged","Customer disputes quality after delivery","Refund is requested or Paystack dispute occurs"]
table(["Scenario","Required system/staff action","Money outcome","Customer message"],[[x,"","",""] for x in cases],[3600,3000,1600,1740],7.4)

pb(); h1("8. Payment, cancellation and refund rules","Paystack will handle Mobile Money/payment credentials. Vikipat must define when money is collected and when production may begin.")
for x in ["Full payment before the order is accepted","Deposit before artwork/design; balance before production","Deposit before production; balance before pickup/delivery","Staff-confirmed orders pay only after final price approval","Additional payment link when scope increases","Partial refund/credit when scope decreases"]:check(x)
field("Which services use which payment rule?",lines=2)
field("Deposit percentage/minimum","If applicable")
field("When balance becomes due")
field("How long an unpaid order/reservation remains valid")
field("Cancellation charges","Separate design, material already used and production started",2)
field("Refund rules","Eligible reasons, deductions and expected timeline",2)
field("Who approves refunds/price overrides","Role or named person; include limits")
field("Manual payment handling","Whether bank transfer/cash is allowed for online orders")

h1("9. Pickup, delivery and installation","Only define fulfilment choices available to storefront customers from the current branch.")
for x in ["Pickup from Mallam–Gbawe Road, opposite Zen Filling Station","Vikipat-arranged delivery","Customer-arranged courier","On-site installation"]:check(x)
table(["Method/zone","How fee is calculated","Expected time","Size/weight restrictions"],[["Pickup","No fee","",""],["Nearby delivery","","",""],["Accra delivery","","",""],["Outside Accra","","",""],["Installation","Travel + labour + equipment?","",""]],[1900,3600,1900,2540],8)
field("Required delivery details","Phone, landmark, digital address, GPS pin, recipient")
field("When delivery estimate needs staff confirmation")
field("Proof of pickup/delivery","OTP, name, signature, photo, etc.")

pb(); h1("10. Products sold directly from the storefront","Use this only for ready-made items and supplies such as DTF film or branded shirts—not configured print services.")
table(["Product","Variants","Selling price","Stock tracked?","Delivery/pickup rule"],[["","","","",""],["","","","",""],["","","","",""],["","","","",""],["","","","",""]],[2200,2300,1600,1500,2340],8)
field("Stock reservation","Reserve at checkout, payment initiation or successful payment?")
field("Out-of-stock/backorder behaviour")
field("Product returns/exchanges","Especially apparel sizes and opened consumables",2)
field("Bundles/quantity pricing","If any")

h1("11. What Vikipat needs to provide","This is the actual hand-over pack needed to build and test the ordering engine.")
items=["Final list of services available online","One completed service sheet per service/variant with a different formula","Current price lists, quantity tables and material rates","At least five real quoted jobs for each major pricing model","Machine/material size limits and impossible option combinations","Artwork requirements, design charges and proof rules","Turnaround and urgent-order rules","Cancellation, refund, rework and delivery rules","Product/SKU list with prices and stock for ready-made goods","Good photos/examples for services, materials, finishes and products","Paystack business account; keys shared later through a secure channel","Names of staff who will review online orders, produce work and approve money changes"]
for x in items:check(x)
callout("Do not include secrets", "Do not place Paystack secret keys, passwords or bank credentials in this workbook. Those will be configured securely during implementation.",GOLD)

h1("12. Final decisions before the build")
table(["Decision","Answer / approved rule"],[["Which services get instant prices?",""],["Which services require staff confirmation before payment?",""],["Can estimates change after artwork review?",""],["When may production begin?",""],["How are extra charges or reductions settled?",""],["What is the final cancellation/refund rule?",""],["Which delivery methods launch first?",""],["Who owns online order review?",""],["Who approves pricing exceptions/refunds?",""],["What historic jobs will be used as pricing tests?",""]],[4400,5540],8.3)
table(["Approved by","Role","Signature","Date"],[["","Operations / pricing","",""],["","Management","",""]],[2700,2400,3000,1840],8.5)

for t in d.tables:
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
d.core_properties.title="Vikipat Core Online Order and Pricing Discovery Workbook"; d.core_properties.subject="Online service configuration, pricing, fulfilment and edge-case discovery"; d.core_properties.author="Vikipat Platform Project"
d.save(OUT); print(OUT)
